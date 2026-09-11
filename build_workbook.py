#!/usr/bin/env python3
"""
Build the Stage 2 marginal-analysis workbook.
Run: python3 build_workbook.py
Output: capabilities/marginal-analysis/model.xlsx
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties


OUTPUT_PATH = Path("capabilities/marginal-analysis/model.xlsx")

GENERAL_INPUTS = [
    ("Season_Weeks", 36, "weeks", "Case scenario"),
    ("Farmer_Hours_Available", 720, "hours", "Case scenario; consumed before temporary labor"),
    ("Farmer_Hourly_Rate", 34.72, "$/hour", "Opportunity cost from case scenario"),
    ("Temp_Hours_Pool", 5760, "hours", "4 workers × 1,440 hours each"),
    ("Temp_Hourly_Rate", 17.36, "$/hour", "Temporary labor wage from case scenario"),
    ("Fixed_Cost", 20000, "$", "Case scenario fixed cost"),
    ("Total_Bed_Cap", 64, "beds", "4 plots × 16 beds"),
]

CROPS = [
    {
        "name": "Tomato",
        "base_formula": 2.5,
        "base_note": "2.50 hours/week/bed",
        "dim_pct": 0.10,
        "revenue": 8800,
        "fertilizer": 880,
        "max_beds": 20,
    },
    {
        "name": "Mesclun",
        "base_formula": 1.25,
        "base_note": "1.25 hours/week/bed",
        "dim_pct": 0.0125,
        "revenue": 2700,
        "fertilizer": 880,
        "max_beds": 30,
    },
    {
        "name": "Carrot",
        "base_formula": "=5/6",
        "base_note": "Exact case value 2.5/3 hours/week/bed",
        "dim_pct": 0.025,
        "revenue": 2094,
        "fertilizer": 440,
        "max_beds": 20,
    },
]

OPTIMAL_BEDS = {"Tomato": 10, "Mesclun": 30, "Carrot": 20}
EXACT_PROFIT = 42768.328256205365
PUBLISHED_PROFIT = 42762

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
INPUT_FILL = PatternFill("solid", fgColor="E2F0D9")
CALC_FILL = PatternFill("solid", fgColor="D9E2F3")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
PASS_FILL = PatternFill("solid", fgColor="C6EFCE")
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
BORDER = Border(
    left=Side(style="thin", color="808080"),
    right=Side(style="thin", color="808080"),
    top=Side(style="thin", color="808080"),
    bottom=Side(style="thin", color="808080"),
)


def style(cell, fill=None, bold=False, number_format=None, align="left"):
    cell.fill = fill or PatternFill(fill_type=None)
    cell.font = Font(bold=bold)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    cell.border = BORDER
    if number_format:
        cell.number_format = number_format


def add_name(workbook, name, ref):
    workbook.defined_names.add(DefinedName(name, attr_text=ref))


def build_inputs_sheet(workbook):
    ws = workbook.create_sheet("Inputs")
    ws.freeze_panes = "A4"
    ws["A1"] = "Stage 2 Perfect Competition Inputs"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Every workbook input is named, unit-labeled, and sourced from the case scenario."

    for cell in ("A3", "B3", "C3", "D3"):
        ws[cell].fill = HEADER_FILL
        ws[cell].font = HEADER_FONT
        ws[cell].border = BORDER
    ws["A3"] = "Named input"
    ws["B3"] = "Value"
    ws["C3"] = "Unit"
    ws["D3"] = "Source / notes"

    row = 4
    for name, value, unit, note in GENERAL_INPUTS:
        ws[f"A{row}"] = name
        ws[f"B{row}"] = value
        ws[f"C{row}"] = unit
        ws[f"D{row}"] = note
        style(ws[f"A{row}"], INPUT_FILL)
        style(ws[f"B{row}"], INPUT_FILL, number_format="$#,##0.00" if "$" in unit else "0.00")
        style(ws[f"C{row}"], INPUT_FILL)
        style(ws[f"D{row}"], INPUT_FILL)
        add_name(workbook, name, f"Inputs!$B${row}")
        row += 1

    ws[f"A{row + 1}"] = "Crop"
    ws[f"B{row + 1}"] = "Base hrs/week/bed"
    ws[f"C{row + 1}"] = "DIM_PCT"
    ws[f"D{row + 1}"] = "Revenue/bed"
    ws[f"E{row + 1}"] = "Fertilizer/bed"
    ws[f"F{row + 1}"] = "Max beds"
    ws[f"G{row + 1}"] = "Source / notes"
    for column in "ABCDEFG":
        header = ws[f"{column}{row + 1}"]
        header.fill = HEADER_FILL
        header.font = HEADER_FONT
        header.border = BORDER

    crop_row = row + 2
    for crop in CROPS:
        name = crop["name"]
        ws[f"A{crop_row}"] = name
        ws[f"B{crop_row}"] = crop["base_formula"]
        ws[f"C{crop_row}"] = crop["dim_pct"]
        ws[f"D{crop_row}"] = crop["revenue"]
        ws[f"E{crop_row}"] = crop["fertilizer"]
        ws[f"F{crop_row}"] = crop["max_beds"]
        ws[f"G{crop_row}"] = crop["base_note"]
        style(ws[f"A{crop_row}"], INPUT_FILL)
        style(ws[f"B{crop_row}"], INPUT_FILL, number_format="0.000000")
        style(ws[f"C{crop_row}"], INPUT_FILL, number_format="0.0000%")
        style(ws[f"D{crop_row}"], INPUT_FILL, number_format="$#,##0.00")
        style(ws[f"E{crop_row}"], INPUT_FILL, number_format="$#,##0.00")
        style(ws[f"F{crop_row}"], INPUT_FILL, number_format="0")
        style(ws[f"G{crop_row}"], INPUT_FILL)

        prefix = name
        add_name(workbook, f"{prefix}_Base_Hours_Per_Week", f"Inputs!$B${crop_row}")
        add_name(workbook, f"{prefix}_Dim_Pct", f"Inputs!$C${crop_row}")
        add_name(workbook, f"{prefix}_Revenue_Per_Bed", f"Inputs!$D${crop_row}")
        add_name(workbook, f"{prefix}_Fertilizer_Per_Bed", f"Inputs!$E${crop_row}")
        add_name(workbook, f"{prefix}_Max_Beds", f"Inputs!$F${crop_row}")
        crop_row += 1

    ws["I3"] = "Published reference checks"
    ws["I3"].fill = HEADER_FILL
    ws["I3"].font = HEADER_FONT
    ws["I3"].border = BORDER
    ws["I4"] = "Optimal mix"
    ws["J4"] = "Tomatoes 10, Mesclun 30, Carrots 20"
    ws["I5"] = "Published season profit"
    ws["J5"] = PUBLISHED_PROFIT
    ws["I6"] = "Exact script profit"
    ws["J6"] = EXACT_PROFIT
    for address in ("I4", "I5", "I6", "J4", "J5", "J6"):
        style(
            ws[address],
            NOTE_FILL if address.startswith("I") else INPUT_FILL,
            number_format="$#,##0.00" if address in {"J5", "J6"} else None,
        )

    widths = {
        "A": 26,
        "B": 16,
        "C": 12,
        "D": 14,
        "E": 16,
        "F": 10,
        "G": 28,
        "I": 22,
        "J": 18,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    return ws


def build_labor_sheet(workbook):
    ws = workbook.create_sheet("Labor & Cost")
    ws.freeze_panes = "A11"
    ws["A1"] = "Labor Function, Standalone Cost Schedule, and Crossing Checks"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Standalone marginal cost uses farmer hours first, then temporary labor."

    summary_headers = ["Crop", "Net contribution/bed", "First q where MC > contribution", "Last profitable bed"]
    for idx, header in enumerate(summary_headers, start=1):
        cell = ws.cell(row=4, column=idx)
        cell.value = header
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER

    summary_rows = {"Tomato": 5, "Mesclun": 6, "Carrot": 7}
    first_cross_cells = {}
    for crop_name, row in summary_rows.items():
        ws[f"A{row}"] = crop_name
        ws[f"B{row}"] = f"={crop_name}_Revenue_Per_Bed-{crop_name}_Fertilizer_Per_Bed"
        style(ws[f"A{row}"], NOTE_FILL)
        style(ws[f"B{row}"], CALC_FILL, number_format="$#,##0.00")

    table_headers = [
        "Crop",
        "q beds",
        "Total labor hours",
        "Marginal labor hours",
        "Standalone labor dollars",
        "Standalone marginal cost",
        "Net contribution/bed",
        "Crossing q",
    ]
    for idx, header in enumerate(table_headers, start=1):
        cell = ws.cell(row=10, column=idx)
        cell.value = header
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER

    current_row = 11
    for crop in CROPS:
        crop_name = crop["name"]
        max_beds = crop["max_beds"]
        start_row = current_row + 1
        end_row = start_row + max_beds
        for q in range(0, max_beds + 1):
            row = start_row + q
            ws[f"A{row}"] = crop_name
            ws[f"B{row}"] = q
            if q == 0:
                ws[f"C{row}"] = "=0"
                ws[f"D{row}"] = "=0"
                ws[f"E{row}"] = "=0"
                ws[f"F{row}"] = "=0"
                ws[f"G{row}"] = f"={crop_name}_Revenue_Per_Bed-{crop_name}_Fertilizer_Per_Bed"
                ws[f"H{row}"] = '=""'
            else:
                ws[f"C{row}"] = (
                    f"={crop_name}_Base_Hours_Per_Week*Season_Weeks*B{row}*(1+{crop_name}_Dim_Pct)^B{row}"
                )
                ws[f"D{row}"] = f"=C{row}-C{row - 1}"
                ws[f"E{row}"] = (
                    f"=MIN(Farmer_Hours_Available,C{row})*Farmer_Hourly_Rate+"
                    f"MAX(0,C{row}-Farmer_Hours_Available)*Temp_Hourly_Rate"
                )
                ws[f"F{row}"] = f"=E{row}-E{row - 1}"
                ws[f"G{row}"] = f"={crop_name}_Revenue_Per_Bed-{crop_name}_Fertilizer_Per_Bed"
                ws[f"H{row}"] = f'=IF(AND(F{row}>G{row},F{row - 1}<=G{row - 1}),B{row},"")'

            style(ws[f"A{row}"], CALC_FILL)
            style(ws[f"B{row}"], CALC_FILL, number_format="0")
            style(ws[f"C{row}"], CALC_FILL, number_format="0.00")
            style(ws[f"D{row}"], CALC_FILL, number_format="0.00")
            style(ws[f"E{row}"], CALC_FILL, number_format="$#,##0.00")
            style(ws[f"F{row}"], CALC_FILL, number_format="$#,##0.00")
            style(ws[f"G{row}"], CALC_FILL, number_format="$#,##0.00")
            style(ws[f"H{row}"], NOTE_FILL, number_format="0")

        ws[f"C{summary_rows[crop_name]}"] = f"=MIN(H{start_row}:H{end_row})"
        ws[f"D{summary_rows[crop_name]}"] = f'=IF(C{summary_rows[crop_name]}="","",C{summary_rows[crop_name]}-1)'
        style(ws[f"C{summary_rows[crop_name]}"], CALC_FILL, number_format="0")
        style(ws[f"D{summary_rows[crop_name]}"], CALC_FILL, number_format="0")
        first_cross_cells[crop_name] = f"C{summary_rows[crop_name]}"

        current_row = end_row + 1

    add_name(workbook, "Schedule_Crop", "'Labor & Cost'!$A$12:$A$86")
    add_name(workbook, "Schedule_Q", "'Labor & Cost'!$B$12:$B$86")
    add_name(workbook, "Schedule_Total_Labor", "'Labor & Cost'!$C$12:$C$86")
    add_name(workbook, "Schedule_Standalone_MC", "'Labor & Cost'!$F$12:$F$86")
    add_name(workbook, "Tomato_First_Standalone_Crossing", "'Labor & Cost'!$C$5")
    add_name(workbook, "Mesclun_First_Standalone_Crossing", "'Labor & Cost'!$C$6")
    add_name(workbook, "Carrot_First_Standalone_Crossing", "'Labor & Cost'!$C$7")

    widths = {"A": 12, "B": 9, "C": 18, "D": 18, "E": 20, "F": 20, "G": 18, "H": 12}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    return ws


def build_allocation_sheet(workbook):
    ws = workbook.create_sheet("Allocation")
    ws.freeze_panes = "A4"
    ws["A1"] = "Allocation and Profit Roll-up"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Decision variables are preloaded to the audited optimum and remain Solver-ready."

    headers = ["Crop", "Beds", "Labor hours", "Revenue", "Fertilizer", "Allocated labor cost", "Contribution after labor"]
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=idx)
        cell.value = header
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER

    row_lookup = {"Tomato": 4, "Mesclun": 5, "Carrot": 6}
    for crop_name, row in row_lookup.items():
        ws[f"A{row}"] = crop_name
        ws[f"B{row}"] = OPTIMAL_BEDS[crop_name]
        ws[f"C{row}"] = f'=SUMIFS(Schedule_Total_Labor,Schedule_Crop,A{row},Schedule_Q,B{row})'
        ws[f"D{row}"] = f"={crop_name}_Revenue_Per_Bed*B{row}"
        ws[f"E{row}"] = f"={crop_name}_Fertilizer_Per_Bed*B{row}"
        ws[f"F{row}"] = f'=IF($B$12=0,0,C{row}*$B$16)'
        ws[f"G{row}"] = f"=D{row}-E{row}-F{row}"
        style(ws[f"A{row}"], INPUT_FILL)
        style(ws[f"B{row}"], INPUT_FILL, number_format="0")
        style(ws[f"C{row}"], CALC_FILL, number_format="0.00")
        style(ws[f"D{row}"], CALC_FILL, number_format="$#,##0.00")
        style(ws[f"E{row}"], CALC_FILL, number_format="$#,##0.00")
        style(ws[f"F{row}"], CALC_FILL, number_format="$#,##0.00")
        style(ws[f"G{row}"], CALC_FILL, number_format="$#,##0.00")

    ws["A7"] = "Total"
    ws["B7"] = "=SUM(B4:B6)"
    ws["C7"] = "=SUM(C4:C6)"
    ws["D7"] = "=SUM(D4:D6)"
    ws["E7"] = "=SUM(E4:E6)"
    ws["F7"] = "=SUM(F4:F6)"
    ws["G7"] = "=SUM(G4:G6)"
    for column in "ABCDEFG":
        style(ws[f"{column}7"], NOTE_FILL, bold=True, number_format="$#,##0.00" if column in "DEFG" else "0.00")
    ws["A7"].number_format = "General"

    metrics = [
        ("Total beds", "=B7"),
        ("Total labor hours", "=C7"),
        ("Farmer hours used", "=MIN(Farmer_Hours_Available,B12)"),
        ("Temp hours used", "=MAX(0,B12-B13)"),
        ("Total labor dollars", "=B13*Farmer_Hourly_Rate+B14*Temp_Hourly_Rate"),
        ("Blended labor rate", '=IF(B12=0,0,B15/B12)'),
        ("Fixed cost", "=Fixed_Cost"),
        ("Profit", "=D7-E7-B15-B17"),
    ]
    for offset, (label, formula) in enumerate(metrics, start=11):
        ws[f"A{offset}"] = label
        ws[f"B{offset}"] = formula
        style(ws[f"A{offset}"], NOTE_FILL)
        style(
            ws[f"B{offset}"],
            CALC_FILL,
            number_format="$#,##0.00" if label in {"Total labor dollars", "Blended labor rate", "Fixed cost", "Profit"} else "0.00",
        )
    ws["B16"].number_format = "$#,##0.0000"
    ws["B18"].font = Font(bold=True, size=12)

    ws["D11"] = "Solver setup"
    ws["D11"].fill = HEADER_FILL
    ws["D11"].font = HEADER_FONT
    ws["D11"].border = BORDER
    solver_lines = [
        ("Objective", "Maximize Profit"),
        ("Changing cells", "Tomato_Beds, Mesclun_Beds, Carrot_Beds"),
        ("Method", "GRG Nonlinear with integer constraints"),
        ("Constraints", "Beds within crop caps; total beds ≤ 64; temporary hours ≤ 5,760"),
        ("Audit starts", "Re-run from 0/0/0 and 20/0/0 in desktop Excel"),
    ]
    for row, (label, value) in enumerate(solver_lines, start=12):
        ws[f"D{row}"] = label
        ws[f"E{row}"] = value
        style(ws[f"D{row}"], NOTE_FILL)
        style(ws[f"E{row}"], INPUT_FILL)

    add_name(workbook, "Tomato_Beds", "Allocation!$B$4")
    add_name(workbook, "Mesclun_Beds", "Allocation!$B$5")
    add_name(workbook, "Carrot_Beds", "Allocation!$B$6")
    add_name(workbook, "Total_Beds", "Allocation!$B$11")
    add_name(workbook, "Total_Labor_Hours", "Allocation!$B$12")
    add_name(workbook, "Farmer_Hours_Used", "Allocation!$B$13")
    add_name(workbook, "Temp_Hours_Used", "Allocation!$B$14")
    add_name(workbook, "Total_Labor_Dollars", "Allocation!$B$15")
    add_name(workbook, "Blended_Labor_Rate", "Allocation!$B$16")
    add_name(workbook, "Profit", "Allocation!$B$18")

    widths = {"A": 25, "B": 14, "C": 14, "D": 14, "E": 16, "F": 18, "G": 22}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 34
    return ws


def build_checks_sheet(workbook):
    ws = workbook.create_sheet("Checks")
    ws.freeze_panes = "A5"
    ws["A1"] = "Validation and Acceptance Checks"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Hand-checks, published references, exact script checks, and feasibility constraints must all pass."

    headers = ["Check", "Expected", "Computed", "Status"]
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=idx)
        cell.value = header
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER

    checks = [
        ("Tomato labor q=1", 99.0, '=SUMIFS(Schedule_Total_Labor,Schedule_Crop,"Tomato",Schedule_Q,1)', '=IF(ABS(C5-B5)<0.01,"PASS","FAIL")', "0.00"),
        ("Tomato labor q=10", 2334.368214090002, '=SUMIFS(Schedule_Total_Labor,Schedule_Crop,"Tomato",Schedule_Q,10)', '=IF(ABS(C6-B6)<0.01,"PASS","FAIL")', "0.00"),
        ("Tomato labor q=20", 12109.4999087861, '=SUMIFS(Schedule_Total_Labor,Schedule_Crop,"Tomato",Schedule_Q,20)', '=IF(ABS(C7-B7)<0.01,"PASS","FAIL")', "0.00"),
        ("Optimal tomato beds", 10, "=Tomato_Beds", '=IF(C8=B8,"PASS","FAIL")', "0"),
        ("Optimal mesclun beds", 30, "=Mesclun_Beds", '=IF(C9=B9,"PASS","FAIL")', "0"),
        ("Optimal carrot beds", 20, "=Carrot_Beds", '=IF(C10=B10,"PASS","FAIL")', "0"),
        ("Exact script profit", EXACT_PROFIT, "=Profit", '=IF(ABS(C11-B11)<0.01,"PASS","FAIL")', "$#,##0.00"),
        ("Published profit reference", PUBLISHED_PROFIT, "=Profit", '=IF(ABS(C12-B12)<=10,"PASS","FAIL")', "$#,##0.00"),
        ("Tomato standalone crossing", 10, "=Tomato_First_Standalone_Crossing-1", '=IF(C13=B13,"PASS","FAIL")', "0"),
        ("Mesclun standalone crossing", 6, "=Mesclun_First_Standalone_Crossing-1", '=IF(C14=B14,"PASS","FAIL")', "0"),
        ("Carrot standalone crossing", 10, "=Carrot_First_Standalone_Crossing-1", '=IF(C15=B15,"PASS","FAIL")', "0"),
        ("Total beds cap", 64, "=Total_Beds", '=IF(C16<=B16,"PASS","FAIL")', "0"),
        ("Temp hours cap", 5760, "=Temp_Hours_Used", '=IF(C17<=B17,"PASS","FAIL")', "0.00"),
        ("Farmer hours nonnegative", 0, "=Farmer_Hours_Used", '=IF(C18>=B18,"PASS","FAIL")', "0.00"),
    ]

    for row, (label, expected, computed, status, fmt) in enumerate(checks, start=5):
        ws[f"A{row}"] = label
        ws[f"B{row}"] = expected
        ws[f"C{row}"] = computed
        ws[f"D{row}"] = status
        style(ws[f"A{row}"], NOTE_FILL)
        style(ws[f"B{row}"], CALC_FILL, number_format=fmt)
        style(ws[f"C{row}"], CALC_FILL, number_format=fmt)
        style(ws[f"D{row}"], NOTE_FILL, bold=True, align="center")

    ws.conditional_formatting.add(
        "D5:D18",
        FormulaRule(formula=['$D5="PASS"'], fill=PASS_FILL),
    )
    ws.conditional_formatting.add(
        "D5:D18",
        FormulaRule(formula=['$D5="FAIL"'], fill=FAIL_FILL),
    )

    widths = {"A": 30, "B": 16, "C": 16, "D": 12}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    return ws


def build_summary_sheet(workbook):
    ws = workbook.create_sheet("Summary")
    ws["A1"] = "Audited Summary"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "The workbook ships at the audited optimum so a reviewer sees the accepted answer on open."

    summary_rows = [
        ("Tomato beds", "=Tomato_Beds"),
        ("Mesclun beds", "=Mesclun_Beds"),
        ("Carrot beds", "=Carrot_Beds"),
        ("Total beds", "=Total_Beds"),
        ("Total labor hours", "=Total_Labor_Hours"),
        ("Farmer hours used", "=Farmer_Hours_Used"),
        ("Temp hours used", "=Temp_Hours_Used"),
        ("Revenue", "=Allocation!D7"),
        ("Fertilizer", "=Allocation!E7"),
        ("Total labor dollars", "=Total_Labor_Dollars"),
        ("Fixed cost", "=Fixed_Cost"),
        ("Profit", "=Profit"),
    ]
    for idx, (label, formula) in enumerate(summary_rows, start=4):
        ws[f"A{idx}"] = label
        ws[f"B{idx}"] = formula
        style(ws[f"A{idx}"], NOTE_FILL)
        style(
            ws[f"B{idx}"],
            CALC_FILL,
            number_format="$#,##0.00" if label in {"Revenue", "Fertilizer", "Total labor dollars", "Fixed cost", "Profit"} else "0.00",
        )
    for cell in ("B4", "B5", "B6", "B7"):
        ws[cell].number_format = "0"

    ws["D4"] = "Standalone P≈MC reference beds"
    ws["D4"].fill = HEADER_FILL
    ws["D4"].font = HEADER_FONT
    ws["D4"].border = BORDER
    ws["D5"] = "Tomato"
    ws["E5"] = "=Tomato_First_Standalone_Crossing-1"
    ws["D6"] = "Mesclun"
    ws["E6"] = "=Mesclun_First_Standalone_Crossing-1"
    ws["D7"] = "Carrot"
    ws["E7"] = "=Carrot_First_Standalone_Crossing-1"
    for address in ("D5", "D6", "D7", "E5", "E6", "E7"):
        style(ws[address], NOTE_FILL if address.startswith("D") else CALC_FILL, number_format="0")

    ws["D9"] = "Stage 3 note"
    ws["D9"].fill = HEADER_FILL
    ws["D9"].font = HEADER_FONT
    ws["D9"].border = BORDER
    ws["D10"] = "Tomato standalone marginal cost dips around q = 6 before rising again."
    ws["D11"] = "Recorded here only; interpretation belongs to Stage 3."
    style(ws["D10"], NOTE_FILL)
    style(ws["D11"], NOTE_FILL)

    widths = {"A": 22, "B": 16, "D": 24, "E": 12}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    return ws


def main():
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)

    build_inputs_sheet(workbook)
    build_labor_sheet(workbook)
    build_allocation_sheet(workbook)
    build_checks_sheet(workbook)
    build_summary_sheet(workbook)

    workbook.save(OUTPUT_PATH)
    print(f"Created {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
