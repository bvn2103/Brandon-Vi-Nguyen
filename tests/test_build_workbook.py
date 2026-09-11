import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import build_workbook


class BuildWorkbookTest(unittest.TestCase):
    def test_builds_required_structure_and_formulas(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "model.xlsx"
            original_output_path = build_workbook.OUTPUT_PATH
            try:
                build_workbook.OUTPUT_PATH = output_path
                build_workbook.main()
            finally:
                build_workbook.OUTPUT_PATH = original_output_path

            workbook = load_workbook(output_path, data_only=False)

            self.assertEqual(
                ["Inputs", "Labor & Cost", "Allocation", "Checks", "Summary"],
                workbook.sheetnames,
            )

            required_names = {
                "Season_Weeks",
                "Farmer_Hours_Available",
                "Farmer_Hourly_Rate",
                "Temp_Hours_Pool",
                "Temp_Hourly_Rate",
                "Fixed_Cost",
                "Tomato_Beds",
                "Mesclun_Beds",
                "Carrot_Beds",
                "Profit",
                "Schedule_Crop",
                "Schedule_Q",
                "Schedule_Total_Labor",
            }
            self.assertTrue(required_names.issubset(set(workbook.defined_names.keys())))

            self.assertEqual("=5/6", workbook["Inputs"]["B15"].value)
            self.assertEqual(
                '=IF(B4=0,0,Tomato_Base_Hours_Per_Week*Season_Weeks*B4*(1+Tomato_Dim_Pct)^B4)',
                workbook["Allocation"]["C4"].value,
            )
            self.assertEqual(
                '=IF(ABS(C11-B11)<0.01,"PASS","FAIL")',
                workbook["Checks"]["D11"].value,
            )
            self.assertEqual(
                '=IF(COUNT(H12:H32)=0,"",SUM(H12:H32))',
                workbook["Labor & Cost"]["C5"].value,
            )
            self.assertEqual(
                '=IF(AND(F23>G23,COUNT(H12:H22)=0),B23,"")',
                workbook["Labor & Cost"]["H23"].value,
            )
            self.assertEqual(
                "=Blended_Labor_Rate",
                workbook["Summary"]["B14"].value,
            )


if __name__ == "__main__":
    unittest.main()
