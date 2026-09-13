#!/usr/bin/env python3
"""Apply professor-requested exactness and audit-check fixes to model workbook.

This is the canonical single script for marginal-analysis workbook corrections.
It updates capabilities/marginal-analysis/model.xlsx to:
- Use exact implied wage formulas (50000/1440 and 25000/1440)
- Keep carrot base hours exact as 2.5/3
- Tighten acceptance checks to cent-level tolerance where values should match
- Correct exact vs published labels/targets

Idempotent and safe to re-run.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
WORKBOOK_PATH = ROOT / "capabilities" / "marginal-analysis" / "model.xlsx"


def set_if_different(cell, value):
    if cell.value != value:
        cell.value = value


def main() -> int:
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(f"Workbook not found: {WORKBOOK_PATH}")

    wb = load_workbook(WORKBOOK_PATH)

    # Inputs sheet
    if "Inputs" in wb.sheetnames:
        ws = wb["Inputs"]
        set_if_different(ws["B5"], "=50000/1440")
        set_if_different(ws["B6"], "=25000/1440")
        set_if_different(ws["B10"], "=2.5/3")

    # Checks sheet
    if "Checks" in wb.sheetnames:
        ws = wb["Checks"]
        set_if_different(ws["A9"], "Exact model target profit")
        set_if_different(ws["B9"], 42761.66)
        set_if_different(ws["A10"], "Published rounded reference")
        set_if_different(ws["B10"], 42762)
        set_if_different(ws["D11"], '=IF(ABS(C11-B11)<=0.01,"PASS","FAIL")')
        set_if_different(ws["D12"], '=IF(ABS(C12-B12)<=0.01,"PASS","FAIL")')

    # Summary sheet (best effort canonical labels)
    if "Summary" in wb.sheetnames:
        ws = wb["Summary"]
        replacements = {
            "A2": "Exact profit (unrounded model):",
            "B2": 42761.66,
            "A3": "Published rounded reference:",
            "B3": 42762,
        }
        for ref, val in replacements.items():
            set_if_different(ws[ref], val)

    wb.save(WORKBOOK_PATH)
    print(f"Updated workbook: {WORKBOOK_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
